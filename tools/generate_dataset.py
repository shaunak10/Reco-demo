#!/usr/bin/env python3
"""
Synthetic dataset generator for the ReconAI demo  (Step 1 of the implementation plan).

Domain: prepaid-recharge aggregator "QuickRecharge Pvt Ltd (SYNTHETIC)" collects customer
payments through a payment gateway ("PayGate Demo") which pays out settlement batches to the
merchant's bank ("Demo Bank").  Three sources must reconcile:

    merchant ledger  --(1:1, ref + gross amount)-->  gateway statement
    gateway statement --(many:1, sum(gross) - sum(fees) == bank credit)--> bank settlement

All data is SYNTHETIC.  Amounts use decimal arithmetic only.

The GROUND TRUTH files are written from the scenario definitions below (the explicit lists of
"special" transactions and batch outcomes), NOT from the output of the matching engine, so the
engine can be scored against them independently.  Every figure is also reproducible by hand:

    Expected_Net  = Gross - Fees_Recorded          (per settlement batch, duplicates excluded)
    Variance      = Bank_Credit - Expected_Net
    Contract fee  = round_half_up(Gross * 1.77%, 2)   (1.5% MDR + 18% GST on MDR)

Run:  python tools/generate_dataset.py          (re-creates /data deterministically)
"""
from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import random

from openpyxl import Workbook

ROOT = Path(__file__).absolute().parent.parent
DATA = ROOT / "data"

CONTRACT_RATE = Decimal("0.0177")
OVERCHARGE_RATE = Decimal("0.0236")
PLANS = [Decimal(x) for x in (149, 199, 239, 299, 349, 479, 599, 719, 999, 1499, 2399, 2999)]
DTH = [Decimal(x) for x in (250, 400, 600, 1000)]
OPERATORS = ["Jio", "Airtel", "Vi", "BSNL"]
DTH_OPERATORS = ["Tata Play", "Dish TV"]
MODES = ["UPI", "UPI", "UPI", "Card", "NetBanking", "Wallet"]


def q2(x: Decimal) -> Decimal:
    return x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def contract_fee(amount: Decimal) -> Decimal:
    return q2(amount * CONTRACT_RATE)


@dataclass
class Txn:
    ref: str
    txn_date: date
    amount: Decimal                       # gateway gross (what was really collected)
    service: str
    operator: str
    mode: str
    fee: Decimal                          # fee recorded by the gateway
    batch: str
    utr: str
    ledger_amount: Decimal | None = None  # defaults to amount
    ledger_ref: str | None = None         # defaults to ref
    gateway_ref: str | None = None        # defaults to ref
    in_ledger: bool = True
    in_gateway: bool = True
    tags: set = field(default_factory=set)


@dataclass
class Batch:
    batch_id: str
    utr: str
    txn_date: date
    count: int
    settle_expected: date
    bank_credit_date: date | None          # None => no credit in the bank statement file
    bank_delta: Decimal = Decimal("0.00")  # bank credit minus expected net
    note: str = ""
    via_api: bool = False                  # credit arrives later through the REST API
    special: dict = field(default_factory=dict)  # txn index -> tag


# --------------------------------------------------------------------------------------
# Scenario definitions (hand-authored)
# --------------------------------------------------------------------------------------
SCENARIO_1 = {
    "name": "scenario1_main",
    "title": "Main demo - 3 days of recharge payments with every exception type",
    "seed": 20261021,
    "batches": [
        Batch("SB-20261016-A", "UTRPG26101601", date(2026, 10, 15), 10, date(2026, 10, 16), date(2026, 10, 16),
              note="Clean batch. One gateway row is repeated in the statement (duplicate).",
              special={3: "GATEWAY_DUPLICATE"}),
        Batch("SB-20261016-B", "UTRPG26101602", date(2026, 10, 15), 9, date(2026, 10, 16), date(2026, 10, 16),
              note="Balances exactly. One txn carries an over-charged fee; one order is booked twice in the ledger.",
              special={2: "LEDGER_DUPLICATE", 5: "FEE_OVERCHARGE"}),
        Batch("SB-20261017-A", "UTRPG26101701", date(2026, 10, 16), 10, date(2026, 10, 17), date(2026, 10, 20),
              note="Bank credited 3 days late but the amount is exact. One txn has a different ledger amount.",
              special={4: "AMOUNT_MISMATCH"}),
        Batch("SB-20261017-B", "UTRPG26101702", date(2026, 10, 16), 8, date(2026, 10, 17), date(2026, 10, 17),
              bank_delta=Decimal("-1200.00"),
              note="Bank credit is short by exactly 1,200.00 - unexplained settlement variance."),
        Batch("SB-20261018-A", "UTRPG26101801", date(2026, 10, 17), 9, date(2026, 10, 18), date(2026, 10, 18),
              note="Balances exactly. Contains a mistyped reference pair and one gateway-only payment.",
              special={2: "REF_TYPO"}),
        Batch("SB-20261018-B", "UTRPG26101802", date(2026, 10, 17), 7, date(2026, 10, 18), None, via_api=True,
              note="No credit in the uploaded statement; the credit arrives late through the bank REST API."),
        Batch("SB-20261019-A", "UTRPG26101901", date(2026, 10, 18), 6, date(2026, 10, 19), None,
              note="Genuinely missing bank credit - never arrives in any source."),
    ],
    "ledger_only": [("RCH-261016-9902", date(2026, 10, 16), Decimal("299.00"), "Mobile Prepaid", "Airtel", "UPI")],
    "gateway_only": [("SB-20261018-A", "RCH-261017-9901", date(2026, 10, 17), Decimal("399.00"), "Mobile Prepaid", "Jio", "Card")],
    "other_bank_credits": [
        (date(2026, 10, 15), "INT261015CR", Decimal("412.50"), "INT.PD:SAVINGS INTEREST CREDIT"),
        (date(2026, 10, 19), "NEFTACME00331", Decimal("10000.00"), "NEFT CR-ACME TRADERS-INVOICE PAYMENT"),
    ],
    "opening_balance": Decimal("250000.00"),
}

SCENARIO_2 = {
    "name": "scenario2_clean",
    "title": "Clean day - everything reconciles, no alert expected",
    "seed": 20261012,
    "batches": [
        Batch("SB-20261013-A", "UTRPG26101301", date(2026, 10, 12), 8, date(2026, 10, 13), date(2026, 10, 13)),
        Batch("SB-20261014-A", "UTRPG26101401", date(2026, 10, 13), 8, date(2026, 10, 14), date(2026, 10, 14)),
        Batch("SB-20261015-A", "UTRPG26101501", date(2026, 10, 14), 8, date(2026, 10, 15), date(2026, 10, 15)),
    ],
    "ledger_only": [],
    "gateway_only": [],
    "other_bank_credits": [],
    "opening_balance": Decimal("120000.00"),
}


def build(spec: dict) -> dict:
    rng = random.Random(spec["seed"])
    txns: list[Txn] = []
    seq_by_day: dict[date, int] = {}

    for b in spec["batches"]:
        for i in range(1, b.count + 1):
            seq_by_day[b.txn_date] = seq_by_day.get(b.txn_date, 0) + 1
            seq = seq_by_day[b.txn_date]
            ref = f"RCH-{b.txn_date:%y%m%d}-{seq:04d}"
            if rng.random() < 0.18:
                amount, service, operator = rng.choice(DTH), "DTH", rng.choice(DTH_OPERATORS)
            else:
                amount, service, operator = rng.choice(PLANS), "Mobile Prepaid", rng.choice(OPERATORS)
            mode = rng.choice(MODES)
            t = Txn(ref, b.txn_date, amount, service, operator, mode, contract_fee(amount), b.batch_id, b.utr)
            tag = b.special.get(i)
            if tag:
                t.tags.add(tag)
            if tag == "FEE_OVERCHARGE":
                t.fee = q2(amount * OVERCHARGE_RATE)
            if tag == "AMOUNT_MISMATCH":
                t.amount = Decimal("569.00")
                t.ledger_amount = Decimal("599.00")
                t.fee = contract_fee(t.amount)
            if tag == "REF_TYPO":
                # gateway shows the letter O where the ledger has the digit zero
                t.gateway_ref = t.ref.replace("-0", "-O", 1)
            txns.append(t)

    for (batch_id, ref, d, amount, service, operator, mode) in spec["gateway_only"]:
        b = next(x for x in spec["batches"] if x.batch_id == batch_id)
        t = Txn(ref, d, amount, service, operator, mode, contract_fee(amount), batch_id, b.utr,
                in_ledger=False, tags={"GATEWAY_ONLY"})
        txns.append(t)
    for (ref, d, amount, service, operator, mode) in spec["ledger_only"]:
        txns.append(Txn(ref, d, amount, service, operator, mode, Decimal("0.00"), "", "",
                        in_gateway=False, tags={"LEDGER_ONLY"}))

    # ------------------------------------------------------------------ ledger rows
    ledger_rows, n = [], 0
    for t in sorted((x for x in txns if x.in_ledger), key=lambda x: (x.txn_date, x.ref)):
        n += 1
        la = t.ledger_amount if t.ledger_amount is not None else t.amount
        ref = t.ledger_ref or t.ref
        ledger_rows.append([f"LDG-{n:06d}", ref, t.txn_date.isoformat(), t.service, t.operator, f"{la:.2f}", t.mode, "SUCCESS"])
        if "LEDGER_DUPLICATE" in t.tags:
            n += 1
            ledger_rows.append([f"LDG-{n:06d}", ref, t.txn_date.isoformat(), t.service, t.operator, f"{la:.2f}", t.mode, "SUCCESS"])

    # ------------------------------------------------------------------ gateway rows
    gateway_rows = []
    gw_seq = 0
    for t in sorted((x for x in txns if x.in_gateway), key=lambda x: (x.txn_date, x.ref)):
        gw_seq += 1
        b = next(x for x in spec["batches"] if x.batch_id == t.batch)
        net = t.amount - t.fee
        row = [t.gateway_ref or t.ref, f"PGT{t.txn_date:%y%m%d}{gw_seq:05d}", t.txn_date, t.amount, t.fee, net,
               b.settle_expected, b.utr, "CAPTURED"]
        gateway_rows.append(row)
        if "GATEWAY_DUPLICATE" in t.tags:
            gateway_rows.append(list(row))

    # ------------------------------------------------------------------ batch maths (independent of engine)
    batch_truth = []
    bank_events = []   # (date, ref, amount, narration)
    for b in spec["batches"]:
        members = [t for t in txns if t.batch == b.batch_id and t.in_gateway]
        gross = sum((t.amount for t in members), Decimal("0.00"))
        fees = sum((t.fee for t in members), Decimal("0.00"))
        expected = gross - fees
        bank_amount = expected + b.bank_delta
        if b.bank_credit_date is not None:
            bank_events.append((b.bank_credit_date, b.utr, bank_amount, f"NEFT CR-PAYGATE DEMO-{b.batch_id}"))
        lag = (b.bank_credit_date - max(t.txn_date for t in members)).days if b.bank_credit_date else None
        if b.bank_credit_date is None and not b.via_api:
            outcome = "MISSING_BANK_CREDIT"
        elif b.bank_credit_date is None and b.via_api:
            outcome = "MISSING_BANK_CREDIT"   # until the API feed is fetched
        elif b.bank_delta != 0:
            outcome = "SETTLEMENT_VARIANCE"
        elif lag is not None and lag > 2:
            outcome = "MATCHED_DELAYED"
        else:
            outcome = "MATCHED"
        after_api = "MATCHED_DELAYED" if b.via_api else outcome
        batch_truth.append({
            "Settlement_Ref": b.utr, "Batch_Id": b.batch_id, "Gateway_Txn_Count_Excl_Duplicates": len(members),
            "Gross": f"{gross:.2f}", "Fees_Recorded": f"{fees:.2f}", "Expected_Net": f"{expected:.2f}",
            "Bank_Credit": f"{bank_amount:.2f}" if b.bank_credit_date else "",
            "Variance_Bank_Minus_Expected": f"{b.bank_delta:.2f}" if b.bank_credit_date else "",
            "Expected_Outcome": outcome, "Expected_After_API_Fetch": after_api, "Note": b.note,
        })
        if b.via_api:
            b._api_amount = bank_amount  # type: ignore[attr-defined]
            b._api_date = date(2026, 10, 21)  # type: ignore[attr-defined]
    for (d, ref, amt, narr) in spec["other_bank_credits"]:
        bank_events.append((d, ref, amt, narr))
    bank_events.sort(key=lambda e: (e[0], e[1]))

    bank_rows, bal = [], spec["opening_balance"]
    for (d, ref, amt, narr) in bank_events:
        bal += amt
        bank_rows.append([d, narr, ref, Decimal("0.00"), amt, bal])

    # ------------------------------------------------------------------ stage-1 ground truth
    s1 = []
    for t in sorted(txns, key=lambda x: (x.txn_date, x.ref)):
        la = t.ledger_amount if t.ledger_amount is not None else t.amount
        if "REF_TYPO" in t.tags:
            s1.append([t.ref, "LEDGER_ONLY", f"{la:.2f}", "", f"{la:.2f}", "Typo pair: ledger ref has digit 0, gateway ref has letter O (no tolerance matching)"])
            s1.append([t.gateway_ref, "GATEWAY_ONLY", "", f"{t.amount:.2f}", f"{-t.amount:.2f}", "Typo pair: counterpart of the ledger-only row above"])
        elif "AMOUNT_MISMATCH" in t.tags:
            s1.append([t.ref, "AMOUNT_MISMATCH", f"{la:.2f}", f"{t.amount:.2f}", f"{la - t.amount:.2f}", "Ledger amount differs from gateway gross"])
        elif "GATEWAY_ONLY" in t.tags:
            s1.append([t.ref, "GATEWAY_ONLY", "", f"{t.amount:.2f}", f"{-t.amount:.2f}", "Payment captured by gateway, no ledger order"])
        elif "LEDGER_ONLY" in t.tags:
            s1.append([t.ref, "LEDGER_ONLY", f"{la:.2f}", "", f"{la:.2f}", "Order booked in ledger, no gateway payment"])
        else:
            s1.append([t.ref, "MATCHED", f"{la:.2f}", f"{t.amount:.2f}", "0.00", ""])
            if "LEDGER_DUPLICATE" in t.tags:
                s1.append([t.ref, "DUPLICATE_LEDGER", f"{la:.2f}", "", "", "Second ledger row with identical ref and amount"])
            if "GATEWAY_DUPLICATE" in t.tags:
                s1.append([t.ref, "DUPLICATE_GATEWAY", "", f"{t.amount:.2f}", "", "Second gateway row with identical ref and amount"])

    # ------------------------------------------------------------------ exception ground truth
    ex = []
    for t in txns:
        la = t.ledger_amount if t.ledger_amount is not None else t.amount
        if "GATEWAY_DUPLICATE" in t.tags:
            ex.append(["DUPLICATE_GATEWAY", t.ref, f"{t.amount:.2f}", "", "Repeated gateway row is excluded from matching and settlement sum"])
        if "LEDGER_DUPLICATE" in t.tags:
            ex.append(["DUPLICATE_LEDGER", t.ref, f"{la:.2f}", "", "Repeated ledger row is excluded from matching"])
        if "AMOUNT_MISMATCH" in t.tags:
            ex.append(["AMOUNT_MISMATCH", t.ref, f"{la - t.amount:.2f}", "", "ledger minus gateway"])
        if "REF_TYPO" in t.tags:
            ex.append(["LEDGER_ONLY", t.ref, f"{la:.2f}", "", "typo pair"])
            ex.append(["GATEWAY_ONLY", t.gateway_ref, f"{t.amount:.2f}", "", "typo pair"])
        if "GATEWAY_ONLY" in t.tags:
            ex.append(["GATEWAY_ONLY", t.ref, f"{t.amount:.2f}", "", "genuine"])
        if "LEDGER_ONLY" in t.tags:
            ex.append(["LEDGER_ONLY", t.ref, f"{la:.2f}", "", "genuine"])
        if "FEE_OVERCHARGE" in t.tags:
            ex.append(["FEE_MISMATCH", t.ref, f"{t.fee - contract_fee(t.amount):.2f}", "", "recorded fee minus contract fee"])
    for row in batch_truth:
        o, ref = row["Expected_Outcome"], row["Settlement_Ref"]
        if o == "MISSING_BANK_CREDIT":
            ex.append(["MISSING_BANK_CREDIT", ref, row["Expected_Net"], "RESOLVED_BY_API" if row["Expected_After_API_Fetch"] != o else "", "expected net settlement with no credit"])
        elif o == "SETTLEMENT_VARIANCE":
            ex.append(["SETTLEMENT_VARIANCE", ref, row["Variance_Bank_Minus_Expected"], "", "bank minus expected net"])
        elif o == "MATCHED_DELAYED":
            ex.append(["DELAYED_SETTLEMENT", ref, "0.00", "", "credit exact but late"])
    for b in spec["batches"]:
        if b.via_api:
            ex.append(["DELAYED_SETTLEMENT", b.utr, "0.00", "AFTER_API_ONLY", "credit exact but late (arrives via API)"])
    for (d, ref, amt, narr) in spec["other_bank_credits"]:
        ex.append(["UNEXPLAINED_BANK_CREDIT", ref, f"{amt:.2f}", "", "bank credit not linked to any gateway settlement"])

    # ------------------------------------------------------------------ API feed (scenario 1 only)
    api_records = []
    for b in spec["batches"]:
        if b.via_api:
            api_records.append({"record_id": f"BNK-20261021-{len(api_records)+1:04d}", "bank_utr": b.utr,
                                "credited_on": b._api_date.isoformat(),  # type: ignore[attr-defined]
                                "credit_amount": f"{b._api_amount:.2f}",  # type: ignore[attr-defined]
                                "remarks": f"NEFT CR-PAYGATE DEMO-{b.batch_id}"})
    if api_records:
        # an overlapping record that is ALSO in the uploaded statement - proves cross-source de-duplication
        ov = next(x for x in bank_events if x[1] == "UTRPG26101801")
        api_records.append({"record_id": f"BNK-20261018-{len(api_records)+1:04d}", "bank_utr": ov[1],
                            "credited_on": ov[0].isoformat(), "credit_amount": f"{ov[2]:.2f}", "remarks": ov[3]})

    return dict(txns=txns, ledger=ledger_rows, gateway=gateway_rows, bank=bank_rows, batches=batch_truth,
                stage1=s1, exceptions=ex, api=api_records)


# --------------------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------------------
def write_csv(path: Path, header: list[str], rows: list[list]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_gateway_xlsx(path: Path, header: list[str], rows: list[list]):
    wb = Workbook()
    ws = wb.active
    ws.title = "PayGate Statement"
    ws.append(header)
    for r in rows:
        ws.append([r[0], r[1], datetime.combine(r[2], datetime.min.time()), float(r[3]), float(r[4]), float(r[5]),
                   datetime.combine(r[6], datetime.min.time()), r[7], r[8]])
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "dd/mm/yyyy"
        row[6].number_format = "dd/mm/yyyy"
    for col, w in zip("ABCDEFGHI", (18, 18, 12, 14, 12, 14, 16, 18, 12)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def fmt_money_bank(d: Decimal) -> str:
    return f"{d:,.2f}"


def write_scenario(spec: dict, alt_headers: bool = False, folder: str | None = None, formats=("csv", "xlsx", "csv")):
    out = DATA / (folder or spec["name"])
    r = build(spec)
    out.mkdir(parents=True, exist_ok=True)

    if not alt_headers:
        lh = ["Ledger_ID", "Order_Ref", "Booking_Date", "Service", "Operator", "Amount_INR", "Pay_Mode", "Order_Status"]
        gh = ["Payment Ref", "Gateway Txn ID", "Txn Date", "Gross Amount", "MDR Fee", "Net Payable", "Settlement Date", "UTR", "Txn Status"]
        bh = ["Credit Date", "Narration", "Bank Reference", "Withdrawal Amt", "Credit Amount", "Closing Balance"]
        names = {"ledger": "merchant_ledger", "gateway": "gateway_statement", "bank": "bank_settlement"}
        gt_map = {
            "ledger": {"Ledger_ID": "IGNORE", "Order_Ref": "TransactionRef", "Booking_Date": "TxnDate", "Service": "IGNORE", "Operator": "IGNORE",
                       "Amount_INR": "Amount", "Pay_Mode": "IGNORE", "Order_Status": "Status"},
            "gateway": {"Payment Ref": "TransactionRef", "Gateway Txn ID": "IGNORE", "Txn Date": "TxnDate", "Gross Amount": "Amount",
                        "MDR Fee": "Fees", "Net Payable": "IGNORE", "Settlement Date": "IGNORE", "UTR": "SettlementRef", "Txn Status": "Status"},
            "bank": {"Credit Date": "TxnDate", "Narration": "IGNORE", "Bank Reference": "SettlementRef", "Withdrawal Amt": "IGNORE",
                     "Credit Amount": "Amount", "Closing Balance": "IGNORE"},
        }
    else:
        lh = ["Internal ID", "Reference No", "Order Date", "Product", "Telecom Operator", "Value (INR)", "Channel", "Order State"]
        gh = ["Merchant Order Id", "PG Transaction Id", "Captured On", "Amount Collected", "Commission", "Amount Payable", "Payout Date", "Bank UTR", "State"]
        bh = ["Value Dt", "Txn Remarks", "Utr No", "Dr Amt", "Cr Amt", "Running Balance"]
        names = {"ledger": "ledger_export", "gateway": "pg_report", "bank": "account_statement"}
        gt_map = {
            "ledger": {"Internal ID": "IGNORE", "Reference No": "TransactionRef", "Order Date": "TxnDate", "Product": "IGNORE",
                       "Telecom Operator": "IGNORE", "Value (INR)": "Amount", "Channel": "IGNORE", "Order State": "Status"},
            "gateway": {"Merchant Order Id": "TransactionRef", "PG Transaction Id": "IGNORE", "Captured On": "TxnDate", "Amount Collected": "Amount",
                        "Commission": "Fees", "Amount Payable": "IGNORE", "Payout Date": "IGNORE", "Bank UTR": "SettlementRef", "State": "Status"},
            "bank": {"Value Dt": "TxnDate", "Txn Remarks": "IGNORE", "Utr No": "SettlementRef", "Dr Amt": "IGNORE", "Cr Amt": "Amount", "Running Balance": "IGNORE"},
        }

    # ledger
    if formats[0] == "csv":
        write_csv(out / f"{names['ledger']}.csv", lh, r["ledger"])
    else:
        wb = Workbook(); ws = wb.active; ws.title = "Ledger"; ws.append(lh)
        for row in r["ledger"]:
            ws.append(row[:5] + [float(row[5])] + row[6:])
        wb.save(out / f"{names['ledger']}.xlsx")
    # gateway
    gw = [row[:] for row in r["gateway"]]
    if formats[1] == "xlsx":
        out.mkdir(parents=True, exist_ok=True)
        write_gateway_xlsx(out / f"{names['gateway']}.xlsx", gh, gw)
    else:
        write_csv(out / f"{names['gateway']}.csv", gh,
                  [[x[0], x[1], x[2].strftime("%d/%m/%Y"), f"{x[3]:.2f}", f"{x[4]:.2f}", f"{x[5]:.2f}", x[6].strftime("%d/%m/%Y"), x[7], x[8]] for x in gw])
    # bank
    bank_rows = [[x[0].strftime("%d-%b-%Y"), x[1], x[2], f"{x[3]:.2f}", fmt_money_bank(x[4]), fmt_money_bank(x[5])] for x in r["bank"]]
    if formats[2] == "csv":
        write_csv(out / f"{names['bank']}.csv", bh, bank_rows)
    else:
        wb = Workbook(); ws = wb.active; ws.title = "Statement"; ws.append(bh)
        for row in r["bank"]:
            ws.append([row[0].strftime("%d-%b-%Y"), row[1], row[2], float(row[3]), float(row[4]), float(row[5])])
        wb.save(out / f"{names['bank']}.xlsx")

    # ground truth (human-checkable)
    write_csv(out / "ground_truth_stage1_ledger_vs_gateway.csv",
              ["Reference", "Expected_Outcome", "Ledger_Amount", "Gateway_Amount", "Ledger_Minus_Gateway", "Note"], r["stage1"])
    write_csv(out / "ground_truth_settlement_batches.csv", list(r["batches"][0].keys()), [list(b.values()) for b in r["batches"]])
    write_csv(out / "ground_truth_exceptions.csv", ["Exception_Type", "Reference", "Exact_Amount", "Flag", "Note"], r["exceptions"])
    write_csv(out / "ground_truth_mapping.csv", ["Source", "Header", "Expected_Canonical_Field"],
              [[src, h, f] for src, m in gt_map.items() for h, f in m.items()])
    if r["api"]:
        (out / "api_bank_credits.json").write_text(json.dumps(r["api"], indent=2), encoding="utf-8")

    manifest = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(out.iterdir()) if p.is_file() and p.name != "MANIFEST.json"}
    (out / "MANIFEST.json").write_text(json.dumps({"scenario": spec["title"], "files_sha256": manifest}, indent=2), encoding="utf-8")
    totals = {
        "ledger_rows": len(r["ledger"]), "gateway_rows": len(r["gateway"]), "bank_rows": len(r["bank"]),
        "gross_total": str(sum((t.amount for t in r["txns"] if t.in_gateway), Decimal("0.00"))),
    }
    print(f"[{out.name}] {totals}")
    return r


def main():
    write_scenario(SCENARIO_1, folder="scenario1_main", formats=("csv", "xlsx", "csv"))
    write_scenario(SCENARIO_2, folder="scenario2_clean", formats=("csv", "xlsx", "csv"))
    # Scenario 3: scenario-1 content under completely different headings + other file formats (tests AI mapping).
    write_scenario(SCENARIO_1, alt_headers=True, folder="scenario3_changed_headers", formats=("xlsx", "csv", "xlsx"))


if __name__ == "__main__":
    main()
